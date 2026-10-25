"""Widget-based terminal desktop. Optional UI dependencies are listed in requirements-tui.txt."""
import asyncio
from pathlib import Path
import subprocess
import sys
from datetime import datetime

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from rich.text import Text
from textual import on
from textual.app import App, ComposeResult
from textual.binding import Binding
from textual.containers import Horizontal, Vertical, VerticalScroll
from textual.widgets import Button, Checkbox, DataTable, Footer, Header, Input, Label, ProgressBar, RichLog, Select, Static

import tui as core


SEASON_LABELS = [("Весна", "spring"), ("Лето", "summer"), ("Осень", "autumn"), ("Зима", "winter")]
PROFILE_LABELS = [("Claude Code", "claude"), ("Codex", "codex"), ("Команда профиля", "default")]
BACKEND_LABELS = [("WSL 2", "wsl"), ("Docker", "docker"), ("Отдельный Windows-пользователь", "user")]


class SandboxDesktop(App):
    """A lightweight dashboard for existing sandbox actions; OS policy remains in the launchers."""

    TITLE = "AI Sandbox"
    SUB_TITLE = "изолированная рабочая среда"
    CSS = """
    Screen { background: #101522; color: #e5e8ef; }
    Header { background: #182033; color: #e5e8ef; }
    Footer { background: #182033; color: #bbc5d7; }
    #layout { height: 1fr; }
    #sidebar { width: 36; min-width: 32; border-right: solid #34415a; padding: 1 1; }
    #main { width: 1fr; height: 1fr; padding: 0 1; }
    #brand { height: 3; color: #73d6ce; text-style: bold; content-align: center middle; }
    .section { color: #e4bd68; text-style: bold; margin-top: 1; height: 1; }
    .hint { color: #9ba8bc; height: auto; margin-bottom: 1; }
    Label { color: #aab5c8; margin-top: 1; height: 1; }
    Select, Input { width: 1fr; margin-bottom: 1; background: #1c2434; border: tall #33425b; }
    Checkbox { margin: 1 0; }
    .actions { height: auto; }
    Button { width: 1fr; min-width: 12; height: 3; margin: 0 1 1 0; border: none; }
    Button.primary { background: #328b60; }
    Button.warning { background: #865c27; }
    Button.danger { background: #7f354a; }
    #hero { height: 18; border: none; padding: 0; margin-bottom: 1; }
    #phase { height: 3; border: round #c58d40; padding: 0 1; color: #edc777; }
    #dash-row { height: 1fr; min-height: 10; }
    #modules { width: 58%; height: 1fr; border: round #4d9c70; }
    #events-box { width: 1fr; height: 1fr; border: round #bb8a3d; margin-left: 1; padding: 0 1; }
    #events { height: 1fr; color: #c9d2e1; }
    #log-box { height: 9; min-height: 6; border: round #2875a3; margin-top: 1; }
    #log { height: 1fr; padding: 0 1; }
    #progress { height: 1; margin: 0 1 1 1; }
    #sidebar-scroll { height: 1fr; }
    .backend-setting { height: auto; }
    .hidden { display: none; }
    DataTable { height: 1fr; }
    """
    BINDINGS = [
        Binding("1", "press_open", "Проекты", show=False),
        Binding("2", "press_create", "Создать", show=False),
        Binding("3", "press_skills", "Навыки", show=False),
        Binding("4", "press_check", "Проверка", show=False),
        Binding("5", "press_selftest", "Самотест", show=False),
        Binding("6", "press_launch", "Запуск", show=False),
        Binding("7", "press_tests", "Тесты", show=False),
        Binding("8", "press_export", "Отчёт", show=False),
        Binding("0", "quit", "Выход", show=False),
        Binding("q", "quit", "Выход"),
    ]

    def __init__(self, platform, root, backend="bubblewrap", season="auto", preview=False):
        super().__init__()
        self.platform = platform
        self.desktop = core.Desktop(platform, root, backend, season)
        self.preview = preview
        self.frame = 0
        self.events = []
        self._project_names = []
        self._running_action = False
        self.ui_logger = logger.bind(component="sandbox-tui")

    def compose(self) -> ComposeResult:
        yield Header()
        with Horizontal(id="layout"):
            with VerticalScroll(id="sidebar"):
                yield Static("AI SANDBOX  ·  DESKTOP", id="brand")
                yield Static("ШАГ 1 · ВЫБРАТЬ ИЛИ СОЗДАТЬ ПРОЕКТ", classes="section")
                yield Label("Существующий проект")
                yield Select([(f"{i + 1}. {name}", name) for i, name in enumerate(self._project_names)] or
                             [("Пока нет проектов — создайте ниже", "__no_projects__")], value=Select.NULL,
                             allow_blank=True, type_to_search=False, id="project-select")
                yield Static("Выберите проект в списке. [1] обновляет его содержимое.",
                             classes="hint", id="project-help")
                yield Button("↻ [1] Обновить список", id="open", variant="primary")
                yield Label("Новый проект")
                yield Input(placeholder="Имя нового проекта", id="project-name")
                yield Input(placeholder="Папка-источник (необязательно)", id="source-path")
                yield Static("Пустой источник создаст пустой проект. Исходная папка останется нетронутой.",
                             classes="hint")
                yield Button("[2] Создать проект", id="create")
                yield Static("ШАГ 2 · НАСТРОИТЬ АГЕНТА", classes="section")
                if self.platform == "windows":
                    yield Label("Изоляция Windows")
                    yield Select(BACKEND_LABELS, value=self.desktop.backend, allow_blank=False,
                                 type_to_search=False, id="backend")
                else:
                    yield Label("Изоляция")
                    yield Static("Linux namespaces · bubblewrap", classes="hint")
                yield Label("Агент")
                yield Select(PROFILE_LABELS, value="claude", allow_blank=False, type_to_search=False, id="profile")
                yield Label("Доступ к копии")
                yield Select([("Только чтение", "read"), ("Редактирование", "edit")],
                             value="read", allow_blank=False, type_to_search=False, id="access")
                yield Label("Сезон сцены")
                yield Select(SEASON_LABELS, value=self.desktop.season, allow_blank=False,
                             type_to_search=False, id="season")
                if self.platform == "windows":
                    yield Label("Пользователь / сетевой адаптер (режим user)")
                    yield Input(value="AIWorker", placeholder="Имя учётной записи", id="username")
                    yield Input(placeholder="GUID адаптера из списка", id="interface")
                    yield Input(placeholder="Путь к агенту .exe для user", id="agent-exe")
                yield Checkbox("Работать без сети", value=False, id="offline")
                yield Static("ШАГ 3 · ДЕЙСТВИЯ", classes="section")
                yield Button("[3] Установить навыки Claude/Codex", id="skills")
                yield Button("[4] Проверить среду", id="check")
                yield Button("[5] Проверить изоляцию", id="selftest", classes="warning")
                yield Button("[7] Запустить тесты проекта", id="tests")
                yield Button("[8] Экспортировать отчёт XLSX", id="export")
                if self.platform == "windows":
                    with Horizontal(classes="actions"):
                        yield Button("Список адаптеров", id="interfaces")
                        yield Button("Настроить user", id="setup-user", classes="warning")
                with Horizontal(classes="actions"):
                    yield Button("[6] ▶ Запустить агента", id="launch", classes="primary")
                    yield Button("Выход", id="exit", classes="danger")
                yield Static("Цифра в скобках запускает действие · q — выход", classes="hint")
            with Vertical(id="main"):
                yield Static("", id="hero")
                yield Static("Фаза  ·  готов к запуску", id="phase")
                with Horizontal(id="dash-row"):
                    yield DataTable(id="modules", zebra_stripes=True)
                    with Vertical(id="events-box"):
                        yield Static("ПРОЕКТ И СЛЕДУЮЩИЙ ШАГ", classes="section")
                        yield Static("Песочница готова. Запуск проверки защиты — по кнопке слева.", id="events")
                with Vertical(id="log-box"):
                    yield Static("ЖУРНАЛ ДЕЙСТВИЙ", classes="section")
                    yield RichLog(id="log", wrap=True, highlight=False, markup=False, max_lines=500)
                yield ProgressBar(total=100, show_eta=False, id="progress")
        yield Footer()

    async def on_mount(self):
        if not self.preview:
            log_dir = self.desktop.root / "logs"
            log_dir.mkdir(parents=True, exist_ok=True)
            logger.remove()
            logger.add(str(log_dir / "tui.log"), rotation="1 MB", retention="7 days",
                       format="{time:YYYY-MM-DD HH:mm:ss} | {level} | {message}",
                       level="INFO", colorize=False, enqueue=False)
            self.ui_logger.info("TUI запущен; platform={platform}, backend={backend}",
                                platform=self.platform, backend=self.desktop.backend)
        table = self.query_one("#modules", DataTable)
        table.add_columns("Компонент", "Состояние", "Проверка")
        table.cursor_type = "row"
        if self.preview:
            self.desktop.items = core.preview_checks()
            self.desktop.healthy = True
            self.desktop.project = "demo"
            self.write_log("ДЕМО: команды и проверки не запускались.", "yellow")
            for button in self.query(Button):
                if button.id not in ("exit",):
                    button.disabled = True
        else:
            self.desktop.items = list(core.checks(self.desktop.root, self.platform))
            self.desktop.healthy = not any(status == "bad" for _, status, _ in self.desktop.items)
            await self.refresh_projects(silent=True)
            self.write_log("Комплект проверен локально. Изоляция ОС запускается отдельно через самотест.", "cyan")
        self.update_module_table()
        self.query_one("#progress", ProgressBar).update(progress=100)
        self.update_summary()
        self.refresh_scene()
        self.set_interval(0.55, self.animate_scene)
        if not self.desktop.healthy:
            self.set_phase("Ошибка в комплекте: запуск отключён до исправления.", "red")

    def write_log(self, message, style=None):
        text = Text(str(message), style=style) if style else Text(str(message))
        self.query_one("#log", RichLog).write(text)

    def set_phase(self, message, color="yellow"):
        self.query_one("#phase", Static).update(Text("Фаза  ·  " + message, style=color))
        self.query_one("#events", Static).update(Text(message))
        self.events.append(message)
        self.events = self.events[-6:]
        self.ui_logger.info("{}", message)

    def update_module_table(self):
        table = self.query_one("#modules", DataTable)
        table.clear(columns=False)
        for name, status, detail in self.desktop.items:
            color = {"ok": "green", "bad": "red", "warn": "yellow"}[status]
            table.add_row(name, Text(status.upper(), style=color), detail)

    def update_summary(self):
        self.update_project_hint()
        if self.desktop.project:
            summary = (
                f"Выбран проект: {self.desktop.project}\n"
                f"Агент: {self.desktop.profile} · доступ: {self.desktop.access}\n"
                f"Изоляция: {self.desktop.backend} · сеть: {'выключена' if self.is_offline() else 'по профилю'}\n"
                f"Комплект: {'готов' if self.desktop.healthy else 'обнаружены ошибки'}\n"
                "Дальше: настройте доступ слева и нажмите [6]."
            )
        else:
            summary = (
                "Существующие проекты находятся в списке слева, вверху.\n"
                "Выберите существующий проект там или нажмите [1], чтобы обновить список.\n"
                "Новый проект: введите имя ниже и нажмите [2].\n"
                "Папка-источник необязательна; оригинал не изменяется."
            )
        self.query_one("#events", Static).update(Text(summary, style="#d6deeb"))

    def update_project_hint(self):
        if self.preview:
            message = "Демо-проект: он показан только для примера, запуск действий отключён."
        elif self.desktop.project in self._project_names:
            message = f"Выбран проект «{self.desktop.project}». Действия ниже применятся к нему."
        elif self._project_names:
            message = f"Найдено проектов: {len(self._project_names)}. Выберите нужный в этом списке."
        else:
            message = "Список пуст. Создайте новый проект ниже или укажите папку для копирования."
        self.query_one("#project-help", Static).update(message)

    def is_offline(self):
        return self.query_one("#offline", Checkbox).value

    def refresh_scene(self):
        widget = self.query_one("#hero", Static)
        width = max(20, widget.size.width)
        height = max(10, widget.size.height)
        pixels, colors = core.seasonal_pixels(self.desktop.season, width, height * 2, self.frame)
        art = Text()
        for y in range(height):
            for x in range(width):
                top = "#%02x%02x%02x" % colors[pixels[y * 2][x]]
                bottom = "#%02x%02x%02x" % colors[pixels[y * 2 + 1][x]]
                art.append("▀", style=f"{top} on {bottom}")
            if y < height - 1:
                art.append("\n")
        widget.update(art)

    def animate_scene(self):
        self.frame += 1
        self.refresh_scene()

    async def refresh_projects(self, silent=False):
        try:
            names = await asyncio.to_thread(self.desktop.names)
            self._project_names = names
            options = [(f"{i + 1}. {name}", name) for i, name in enumerate(names)] or [
                ("Пока нет проектов — создайте ниже", "__no_projects__")]
            select = self.query_one("#project-select", Select)
            select.set_options(options)
            if self.desktop.project in names:
                select.value = self.desktop.project
            else:
                self.desktop.project = ""
                select.value = Select.NULL
            self.update_project_hint()
            if not silent:
                self.set_phase(f"Список обновлён: {len(names)} проектов. Выбери проект в списке слева.", "green")
        except (OSError, ValueError, subprocess.TimeoutExpired) as exc:
            if not silent:
                self.write_log(f"Не удалось получить список проектов: {exc}", "red")
                self.set_phase("Ошибка чтения списка проектов", "red")

    @on(Select.Changed)
    async def selection_changed(self, event: Select.Changed):
        value = event.value
        if event.select.id == "profile" and value:
            self.desktop.profile = str(value)
        elif event.select.id == "access" and value:
            self.desktop.access = str(value)
        elif event.select.id == "season" and value:
            self.desktop.season = str(value)
            self.frame = 0
            self.refresh_scene()
        elif event.select.id == "project-select" and value:
            if str(value) in self._project_names:
                self.desktop.project = str(value)
                self.update_project_hint()
        elif event.select.id == "backend" and value:
            self.desktop.backend = str(value)
            self.desktop.project = ""
            await self.refresh_projects(silent=True)
        self.update_summary()

    @on(Button.Pressed)
    def button_pressed(self, event: Button.Pressed):
        action = event.button.id
        if action == "exit":
            self.exit()
        elif action:
            self.run_worker(self.handle_action(action), name=f"sandbox-{action}", group="actions", exclusive=True)

    def action_press(self, action):
        if action == "quit":
            self.exit()
        else:
            self.run_worker(self.handle_action(action), name=f"sandbox-{action}", group="actions", exclusive=True)

    def action_press_open(self): self.action_press("open")
    def action_press_create(self): self.action_press("create")
    def action_press_skills(self): self.action_press("skills")
    def action_press_check(self): self.action_press("check")
    def action_press_selftest(self): self.action_press("selftest")
    def action_press_launch(self): self.action_press("launch")
    def action_press_tests(self): self.action_press("tests")
    def action_press_export(self): self.action_press("export")

    async def handle_action(self, action):
        if self.preview:
            return
        if action == "open":
            await self.refresh_projects()
            if self.desktop.project:
                self.set_phase(f"Выбран проект: {self.desktop.project}", "green")
            else:
                self.set_phase("Список обновлён. Выберите проект выше или создайте новый по кнопке [2].")
            self.update_summary()
            self.refresh_scene()
            return
        if action == "export":
            self.export_report()
            return
        if action == "create":
            name = self.query_one("#project-name", Input).value.strip()
            source = self.query_one("#source-path", Input).value.strip()
            if not name:
                self.set_phase("Введите имя нового проекта.", "red")
                return
            extra = ["--copy", source] if source else []
            result = await self.run_quiet("new", name, extra=extra)
            if result == 0:
                self.desktop.project = name
                await self.refresh_projects(silent=True)
                await self.run_quiet("skills", name)
                self.write_log("Навыки Claude и Codex подготовлены в копии проекта.", "green")
                self.query_one("#project-name", Input).value = ""
                self.set_phase(f"Создан проект: {name}", "green")
            self.update_summary()
            self.refresh_scene()
            return
        if action in ("interfaces", "setup-user"):
            if self.platform != "windows" or self.desktop.backend != "user":
                self.set_phase("Выберите режим Windows user.", "red")
                return
            self.desktop.username = self.query_one("#username", Input).value.strip() or "AIWorker"
            self.desktop.interface = self.query_one("#interface", Input).value.strip()
            if action == "setup-user" and not self.desktop.interface:
                self.set_phase("Введите GUID адаптера из списка сетевых интерфейсов.", "red")
                return
            result = await self.run_quiet(action)
            self.set_phase("Команда завершена" if result == 0 else f"Ошибка команды, код {result}",
                           "green" if result == 0 else "red")
            return
        if action == "skills":
            if not await self.require_project(): return
            result = await self.run_quiet("skills", self.desktop.project)
            self.set_phase("Навыки установлены" if result == 0 else f"Ошибка навыков, код {result}",
                           "green" if result == 0 else "red")
            return
        if action in ("check", "selftest"):
            if action == "selftest" and not self.desktop.healthy:
                self.set_phase("Сначала исправьте ошибку комплекта.", "red")
                return
            args = (self.desktop.profile,) if action == "selftest" else ()
            result = await self.run_quiet(action, *args)
            self.set_phase(("Проверка завершена" if result == 0 else f"Проверка завершилась с кодом {result}"),
                           "green" if result == 0 else "red")
            return
        if action in ("launch", "tests"):
            if not self.desktop.healthy:
                self.set_phase("Запуск заблокирован: комплект не прошёл проверку.", "red")
                return
            if not await self.require_project(): return
            if action == "tests":
                command = self.detect_tests()
                if not command:
                    self.set_phase("В копии не найден стандартный тестовый набор.", "yellow")
                    return
                extra = (["--offline"] if self.platform == "windows" else ["--net", "off"])
                argv = self.desktop.argv("run", "default", self.desktop.project, extra=extra + ["--"] + command)
            else:
                extra = ["--access", self.desktop.access]
                if self.is_offline():
                    extra += ["--offline"] if self.platform == "windows" else ["--net", "off"]
                if self.platform == "windows" and self.desktop.backend == "user":
                    exe = self.query_one("#agent-exe", Input).value.strip()
                    if not exe:
                        self.set_phase("Для режима user укажите путь к .exe агента.", "red")
                        return
                    extra += ["--", exe]
                argv = self.desktop.argv("run", self.desktop.profile, self.desktop.project, extra=extra)
            await self.run_interactive(argv)

    def detect_tests(self):
        if self.platform == "windows" and self.desktop.backend == "user":
            return None
        folder = None if self.platform == "windows" and self.desktop.backend == "wsl" else self.desktop.workspace_root() / self.desktop.project
        for marker, command in (("tests", ["python3", "-m", "unittest", "discover", "-s", "tests", "-v"]),
                                ("package.json", ["npm", "test"]), ("Cargo.toml", ["cargo", "test"]),
                                ("go.mod", ["go", "test", "./..."])):
            if folder is None or ((folder / marker).exists() and not (folder / marker).is_symlink()):
                return command
        return None

    def export_report(self):
        """Export the local module report and current UI settings without collecting host data."""
        reports = self.desktop.root / "state" / "reports"
        reports.mkdir(parents=True, exist_ok=True)
        path = reports / ("sandbox-report-" + datetime.now().strftime("%Y%m%d-%H%M%S") + ".xlsx")
        book = Workbook()
        modules = book.active
        modules.title = "Проверки комплекта"
        modules.append(["Компонент", "Статус", "Проверка"])
        for item in self.desktop.items:
            modules.append(list(item))
        summary = book.create_sheet("Сводка")
        summary.append(["Параметр", "Значение"])
        for key, value in (("Платформа", self.platform), ("Режим изоляции", self.desktop.backend),
                           ("Проект", self.desktop.project or "не выбран"), ("Профиль", self.desktop.profile),
                           ("Доступ", self.desktop.access), ("Сезон", core.SEASONS[self.desktop.season]),
                           ("Состояние комплекта", "готов" if self.desktop.healthy else "есть ошибки")):
            summary.append([key, value])
        for sheet in (modules, summary):
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="263B59")
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(52, max(14, max(len(str(c.value or "")) for c in column) + 2))
        table = Table(displayName="ModuleChecks", ref=modules.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True,
                                              showColumnStripes=False)
        modules.add_table(table)
        book.save(path)
        self.write_log(f"Отчёт сохранён: {path}", "green")
        self.set_phase(f"Отчёт XLSX сохранён: {path.name}", "green")

    async def require_project(self):
        await self.refresh_projects(silent=True)
        if self.desktop.project and self.desktop.project in self._project_names:
            return True
        self.set_phase("Сначала выберите проект в списке слева.", "red")
        return False

    async def run_quiet(self, action, *positional, extra=()):
        argv = self.desktop.argv(action, *positional, extra=extra)
        self.write_log("$ " + " ".join(argv), "#9ba8bc")
        self.set_phase(f"Выполняется {action}…")
        try:
            result = await asyncio.to_thread(subprocess.run, argv, capture_output=True, text=True,
                                             encoding="utf-8", errors="replace", check=False)
            for line in (result.stdout + "\n" + result.stderr).splitlines()[-120:]:
                if line.strip(): self.write_log(line)
            return result.returncode
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.write_log(str(exc), "red")
            return 1

    async def run_interactive(self, argv):
        self._running_action = True
        self.set_phase("Рабочий стол приостановлен; управление передано CLI-агенту.", "yellow")
        self.write_log("Запуск дочерней команды с действующей изоляцией ОС.", "cyan")
        try:
            with self.suspend():
                result = await asyncio.to_thread(subprocess.run, argv, check=False)
            self.set_phase(f"Агент завершился с кодом {result.returncode}.", "green" if result.returncode == 0 else "yellow")
            self.write_log(f"Код завершения агента: {result.returncode}", "green" if result.returncode == 0 else "yellow")
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.set_phase(f"Ошибка запуска: {exc}", "red")
            self.write_log(str(exc), "red")
        finally:
            self._running_action = False
            self.refresh_scene()


def main(platform, root, backend="bubblewrap", season="auto", preview=False):
    app = SandboxDesktop(platform, root, backend, season, preview=preview)
    app.run()
    return 0
